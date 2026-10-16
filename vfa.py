
import argparse
import math
import sys
import tempfile
from pathlib import Path
from datetime import datetime, timedelta

import cv2
import numpy as np
import torch

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# SETTINGS
# ============================================================

PROJECT_DIR = Path(__file__).resolve().parent

CONFIDENCE = 0.30
MAX_TRACK_DISTANCE = 100
MAX_MISSED_FRAMES = 30

VEHICLE_CLASSES = {"car", "motorcycle", "bus", "truck"}

EXCEL_COLUMNS = [
    "ENTRY",
    "ENTRY_TIME",
    "EXIT",
    "EXIT_TIME",
    "VEHICLE_ID",
    "COLOUR_RGB",
    "VEHICLE_TYPE",
    "SNAPSHOT",
    "TIME_IN_AREA",
]


# ============================================================
# EXCEL FORMATTING
# ============================================================

def create_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Vehicle Analysis"

    ws.append(EXCEL_COLUMNS)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="17365D",
    )

    thin_grey = Side(style="thin", color="D9E2F3")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(
            name="Arial",
            size=10,
            bold=True,
            color="FFFFFF",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(bottom=thin_grey)

    widths = {
        "A": 14,
        "B": 15,
        "C": 14,
        "D": 15,
        "E": 16,
        "F": 22,
        "G": 18,
        "H": 24,
        "I": 18,
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:I1"
    ws.sheet_view.showGridLines = False

    return wb, ws


def format_excel_row(ws, row_number):
    for col in range(1, len(EXCEL_COLUMNS) + 1):
        cell = ws.cell(row_number, col)

        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(
            bottom=Side(style="thin", color="D9E2F3")
        )

        if row_number % 2 == 0:
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="F2F6FC",
            )

    ws.row_dimensions[row_number].height = 70


def add_snapshot_to_excel(ws, row_number, crop, temporary_folder):
    """
    Temporarily saves the crop so openpyxl can embed it in Excel.
    The temporary image folder is deleted after the workbook is saved.
    No separate snapshot folder is left behind.
    """

    if crop is None or crop.size == 0:
        ws.cell(row_number, 8).value = "Unavailable"
        return

    vehicle_id = str(ws.cell(row_number, 5).value)
    image_path = Path(temporary_folder) / f"{vehicle_id}_{row_number}.jpg"

    if not cv2.imwrite(str(image_path), crop):
        ws.cell(row_number, 8).value = "Unavailable"
        return

    try:
        image = ExcelImage(str(image_path))
        image.width = 160
        image.height = 90

        ws.add_image(image, f"H{row_number}")
        ws.cell(row_number, 8).value = "Vehicle snapshot"

    except Exception as error:
        print(f"Could not embed snapshot for {vehicle_id}: {error}")
        ws.cell(row_number, 8).value = "Unavailable"


# ============================================================
# EXACT RGB COLOUR COUNTING
# ============================================================

def get_exact_rgb_colour(crop):
    """
    Finds the exact RGB triplet that occurs most often in the
    rectangular vehicle crop. No segmentation or colour grouping.
    """

    if crop is None or crop.size == 0:
        return "Unknown"

    rgb = cv2.cvtColor(crop, cv2.COLOR_BGR2RGB)
    pixels = rgb.reshape(-1, 3)

    packed = (
        (pixels[:, 0].astype(np.uint32) << 16)
        | (pixels[:, 1].astype(np.uint32) << 8)
        | pixels[:, 2].astype(np.uint32)
    )

    values, counts = np.unique(packed, return_counts=True)
    colour = int(values[np.argmax(counts)])

    red = (colour >> 16) & 255
    green = (colour >> 8) & 255
    blue = colour & 255

    return f"RGB({red}, {green}, {blue})"


# ============================================================
# GATE CONFIGURATION
# ============================================================

def configure_gate(frame, name):
    """
    Click the two endpoints of the gate line.
    Then click a point on the INSIDE of the monitored area.
    Press ENTER to accept, R to reset, or Q to cancel.
    """

    points = []
    window = f"Configure {name}"

    def mouse_callback(event, x, y, flags, parameter):
        if event == cv2.EVENT_LBUTTONDOWN and len(points) < 3:
            points.append((x, y))

    cv2.namedWindow(window)
    cv2.setMouseCallback(window, mouse_callback)

    gate = None

    while True:
        display = frame.copy()

        instructions = [
            f"{name}: click the TWO endpoints of the gate line",
            "Then click a point INSIDE the monitored area",
            "ENTER = accept | R = reset | Q = cancel",
        ]

        for i, text in enumerate(instructions):
            cv2.putText(
                display,
                text,
                (12, 28 + i * 28),
                cv2.FONT_HERSHEY_SIMPLEX,
                0.62,
                (0, 255, 255),
                2,
            )

        if len(points) >= 1:
            cv2.circle(display, points[0], 6, (0, 255, 0), -1)

        if len(points) >= 2:
            cv2.line(display, points[0], points[1], (255, 0, 255), 3)
            cv2.circle(display, points[1], 6, (0, 255, 0), -1)

        if len(points) >= 3:
            cv2.circle(display, points[2], 8, (0, 0, 255), -1)
            cv2.putText(
                display,
                "INSIDE",
                (points[2][0] + 8, points[2][1]),
                cv2.FONT_HERSHEY_SIMPLEX,
                0.65,
                (0, 0, 255),
                2,
            )

        cv2.imshow(window, display)
        key = cv2.waitKey(20) & 0xFF

        if key == ord("r"):
            points.clear()

        elif key in (10, 13) and len(points) == 3:
            p1, p2, inside_point = points

            if p1 == p2:
                print("The two gate endpoints must be different.")
                points.clear()
                continue

            cross = (
                (p2[0] - p1[0]) * (inside_point[1] - p1[1])
                - (p2[1] - p1[1]) * (inside_point[0] - p1[0])
            )

            if cross == 0:
                print("Click an inside point that is not on the gate.")
                points.clear()
                continue

            gate = {
                "name": name,
                "p1": p1,
                "p2": p2,
                "inside_sign": 1 if cross > 0 else -1,
            }
            break

        elif key == ord("q"):
            break

    cv2.destroyWindow(window)

    if gate is None:
        raise RuntimeError(f"Configuration cancelled for {name}.")

    return gate


def signed_distance(point, gate):
    x, y = point
    x1, y1 = gate["p1"]
    x2, y2 = gate["p2"]

    length = math.hypot(x2 - x1, y2 - y1)

    if length == 0:
        return 0

    cross = (x2 - x1) * (y - y1) - (y2 - y1) * (x - x1)

    return cross / length


def crosses_finite_gate(previous, current, gate):
    """
    Returns True only if movement crosses the actual gate segment,
    not merely the infinite extension of the line.
    """

    x1, y1 = previous
    x2, y2 = current

    gx1, gy1 = gate["p1"]
    gx2, gy2 = gate["p2"]

    dx = x2 - x1
    dy = y2 - y1
    gdx = gx2 - gx1
    gdy = gy2 - gy1

    denominator = dx * gdy - dy * gdx

    if abs(denominator) < 1e-9:
        return False

    qx = gx1 - x1
    qy = gy1 - y1

    movement_fraction = (qx * gdy - qy * gdx) / denominator
    gate_fraction = (qx * dy - qy * dx) / denominator

    return (
        0 <= movement_fraction <= 1
        and 0 <= gate_fraction <= 1
    )


# ============================================================
# VEHICLE TRACKING
# ============================================================

def calculate_iou(a, b):
    ax1, ay1, ax2, ay2 = a
    bx1, by1, bx2, by2 = b

    ix1 = max(ax1, bx1)
    iy1 = max(ay1, by1)
    ix2 = min(ax2, bx2)
    iy2 = min(ay2, by2)

    intersection = max(0, ix2 - ix1) * max(0, iy2 - iy1)

    area_a = max(0, ax2 - ax1) * max(0, ay2 - ay1)
    area_b = max(0, bx2 - bx1) * max(0, by2 - by1)
    union = area_a + area_b - intersection

    return intersection / union if union > 0 else 0


def track_vehicles(detections, tracks, next_id):
    """
    Basic single-camera tracking using bounding-box overlap
    and centroid distance. IDs are displayed as V001, V002, etc.
    """

    candidates = []

    for di, detection in enumerate(detections):
        cx, cy = detection["centroid"]

        for vehicle_id, track in tracks.items():
            tx, ty = track["centroid"]
            distance = math.hypot(cx - tx, cy - ty)
            overlap = calculate_iou(detection["bbox"], track["bbox"])

            if distance <= MAX_TRACK_DISTANCE or overlap >= 0.10:
                score = distance / MAX_TRACK_DISTANCE + (1 - overlap)
                candidates.append((score, di, vehicle_id))

    candidates.sort()

    assigned_detections = set()
    assigned_tracks = set()
    assignments = {}

    for score, di, vehicle_id in candidates:
        if di in assigned_detections or vehicle_id in assigned_tracks:
            continue

        assignments[di] = vehicle_id
        assigned_detections.add(di)
        assigned_tracks.add(vehicle_id)

    for vehicle_id in list(tracks.keys()):
        if vehicle_id not in assigned_tracks:
            tracks[vehicle_id]["missed"] += 1

            if tracks[vehicle_id]["missed"] > MAX_MISSED_FRAMES:
                del tracks[vehicle_id]

    visible_ids = set()

    for di, detection in enumerate(detections):
        if di in assignments:
            vehicle_id = assignments[di]
            track = tracks[vehicle_id]

            track["previous_centroid"] = track["centroid"]
            track["centroid"] = detection["centroid"]
            track["bbox"] = detection["bbox"]
            track["vehicle_type"] = detection["vehicle_type"]
            track["crop"] = detection["crop"]
            track["missed"] = 0

        else:
            vehicle_id = f"V{next_id:03d}"
            next_id += 1

            tracks[vehicle_id] = {
                "centroid": detection["centroid"],
                "previous_centroid": None,
                "bbox": detection["bbox"],
                "vehicle_type": detection["vehicle_type"],
                "crop": detection["crop"],
                "missed": 0,
                "gate_sides": {},
            }

        detection["vehicle_id"] = vehicle_id
        visible_ids.add(vehicle_id)

    return next_id, visible_ids


# ============================================================
# TIME AND ENTRY/EXIT RECORDS
# ============================================================

def timestamp_at(start_time, elapsed_seconds):
    return (start_time + timedelta(seconds=elapsed_seconds)).strftime("%H:%M:%S")


def duration_string(seconds):
    seconds = max(0, int(round(seconds)))
    hours = seconds // 3600
    minutes = (seconds % 3600) // 60
    seconds = seconds % 60

    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


def process_gate_events(
    tracks,
    visible_ids,
    gates,
    active_visits,
    worksheet,
    start_time,
    elapsed_seconds,
    temporary_folder,
):
    for vehicle_id in visible_ids:
        track = tracks[vehicle_id]
        previous = track["previous_centroid"]
        current = track["centroid"]

        if previous is None:
            continue

        for gate in gates:
            name = gate["name"]

            previous_distance = signed_distance(previous, gate)
            current_distance = signed_distance(current, gate)

            previous_side = track["gate_sides"].get(name)
            current_side = 1 if current_distance > 0 else -1 if current_distance < 0 else 0

            if current_side == 0:
                continue

            track["gate_sides"][name] = current_side

            if previous_side is None or previous_side == current_side:
                continue

            if not crosses_finite_gate(previous, current, gate):
                continue

            was_inside = previous_side == gate["inside_sign"]
            is_inside = current_side == gate["inside_sign"]

            # Outside -> inside: ENTRY through either gate.
            if not was_inside and is_inside:
                if vehicle_id in active_visits:
                    continue

                row = worksheet.max_row + 1
                entry_time = timestamp_at(start_time, elapsed_seconds)
                colour = get_exact_rgb_colour(track["crop"])

                worksheet.cell(row, 1).value = name
                worksheet.cell(row, 2).value = entry_time
                worksheet.cell(row, 3).value = ""
                worksheet.cell(row, 4).value = ""
                worksheet.cell(row, 5).value = vehicle_id
                worksheet.cell(row, 6).value = colour
                worksheet.cell(row, 7).value = track["vehicle_type"]
                worksheet.cell(row, 8).value = ""
                worksheet.cell(row, 9).value = ""

                format_excel_row(worksheet, row)
                add_snapshot_to_excel(
                    worksheet,
                    row,
                    track["crop"],
                    temporary_folder,
                )

                active_visits[vehicle_id] = {
                    "row": row,
                    "entry_gate": name,
                    "elapsed": elapsed_seconds,
                }

                print(f"ENTRY  | {vehicle_id} | {name} | {entry_time}")

            # Inside -> outside: EXIT through either gate.
            elif was_inside and not is_inside:
                exit_time = timestamp_at(start_time, elapsed_seconds)

                if vehicle_id in active_visits:
                    visit = active_visits.pop(vehicle_id)
                    row = visit["row"]

                    worksheet.cell(row, 3).value = name
                    worksheet.cell(row, 4).value = exit_time
                    worksheet.cell(row, 9).value = duration_string(
                        elapsed_seconds - visit["elapsed"]
                    )

                    print(
                        f"EXIT   | {vehicle_id} | {name} | {exit_time} "
                        f"| TIME: {worksheet.cell(row, 9).value}"
                    )

                else:
                    # Entry was not observed in this video.
                    row = worksheet.max_row + 1
                    colour = get_exact_rgb_colour(track["crop"])

                    worksheet.cell(row, 1).value = ""
                    worksheet.cell(row, 2).value = ""
                    worksheet.cell(row, 3).value = name
                    worksheet.cell(row, 4).value = exit_time
                    worksheet.cell(row, 5).value = vehicle_id
                    worksheet.cell(row, 6).value = colour
                    worksheet.cell(row, 7).value = track["vehicle_type"]
                    worksheet.cell(row, 8).value = ""
                    worksheet.cell(row, 9).value = ""

                    format_excel_row(worksheet, row)
                    add_snapshot_to_excel(
                        worksheet,
                        row,
                        track["crop"],
                        temporary_folder,
                    )

                    print(
                        f"EXIT ONLY | {vehicle_id} | {name} | {exit_time}"
                    )


# ============================================================
# MAIN
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description="Single-camera, two-gate vehicle analysis"
    )
    parser.add_argument("--video", required=True, help="Input video filename")
    parser.add_argument(
        "--output-xlsx",
        default="vehicle_results.xlsx",
        help="Output Excel filename",
    )

    args = parser.parse_args()

    video_path = Path(args.video).expanduser().resolve()
    output_path = Path(args.output_xlsx).expanduser().resolve()

    if not video_path.exists():
        print(f"ERROR: Video not found: {video_path}")
        sys.exit(1)

    output_path.parent.mkdir(parents=True, exist_ok=True)

    print("\nLoading YOLOv5...")

    try:
        model = torch.hub.load(
            str(PROJECT_DIR),
            "yolov5s",
            source="local",
        )
    except Exception as error:
        print(f"Could not load the local YOLOv5 model:\n{error}")
        sys.exit(1)

    model.conf = CONFIDENCE
    model.iou = 0.45

    capture = cv2.VideoCapture(str(video_path))

    if not capture.isOpened():
        print(f"ERROR: Could not open video: {video_path}")
        sys.exit(1)

    fps = capture.get(cv2.CAP_PROP_FPS)
    if not fps or fps <= 0:
        fps = 25.0

    total_frames = int(capture.get(cv2.CAP_PROP_FRAME_COUNT))

    ok, first_frame = capture.read()
    if not ok:
        capture.release()
        print("ERROR: Could not read the video.")
        sys.exit(1)

    print("\nGATE SETUP")
    print("For each gate, click the two endpoints, then click INSIDE the area.")
    print("The same gate can be used for entry and exit.")

    try:
        gate_a = configure_gate(first_frame, "Gate A")
        gate_b = configure_gate(first_frame, "Gate B")
    except RuntimeError as error:
        capture.release()
        cv2.destroyAllWindows()
        print(error)
        sys.exit(1)

    while True:
        start_input = input(
            "\nEnter CCTV start time as HH:MM:SS, or press ENTER for 00:00:00: "
        ).strip()

        if not start_input:
            start_input = "00:00:00"

        try:
            start_time = datetime.strptime(start_input, "%H:%M:%S")
            break
        except ValueError:
            print("Invalid time. Example: 10:30:00")

    workbook, worksheet = create_excel()

    tracks = {}
    active_visits = {}
    next_id = 1
    frame_number = 0

    # Temporary images are used only to embed snapshots in Excel.
    # The temporary folder is removed after the workbook is saved.
    with tempfile.TemporaryDirectory(prefix="vfas_temp_") as temporary_folder:
        print("\nStarting vehicle analysis.")
        print("Press Q in the video preview to stop early.\n")

        capture.set(cv2.CAP_PROP_POS_FRAMES, 0)

        try:
            while True:
                ok, frame = capture.read()
                if not ok:
                    break

                frame_number += 1
                elapsed_seconds = (frame_number - 1) / fps

                results = model(frame)
                predictions = results.xyxy[0].cpu().numpy()

                height, width = frame.shape[:2]
                detections = []

                for prediction in predictions:
                    x1, y1, x2, y2, confidence, class_id = prediction
                    class_id = int(class_id)
                    vehicle_type = str(model.names[class_id]).lower()

                    if vehicle_type not in VEHICLE_CLASSES:
                        continue

                    if confidence < CONFIDENCE:
                        continue

                    x1 = max(0, min(width - 1, int(x1)))
                    y1 = max(0, min(height - 1, int(y1)))
                    x2 = max(0, min(width, int(x2)))
                    y2 = max(0, min(height, int(y2)))

                    if x2 <= x1 or y2 <= y1:
                        continue

                    crop = frame[y1:y2, x1:x2].copy()

                    detections.append({
                        "bbox": (x1, y1, x2, y2),
                        "centroid": ((x1 + x2) // 2, (y1 + y2) // 2),
                        "vehicle_type": vehicle_type,
                        "crop": crop,
                    })

                next_id, visible_ids = track_vehicles(
                    detections,
                    tracks,
                    next_id,
                )

                process_gate_events(
                    tracks=tracks,
                    visible_ids=visible_ids,
                    gates=[gate_a, gate_b],
                    active_visits=active_visits,
                    worksheet=worksheet,
                    start_time=start_time,
                    elapsed_seconds=elapsed_seconds,
                    temporary_folder=temporary_folder,
                )

                # Draw gates on the video preview.
                for gate in (gate_a, gate_b):
                    colour = (0, 255, 0) if gate["name"] == "Gate A" else (255, 0, 255)

                    cv2.line(
                        frame,
                        gate["p1"],
                        gate["p2"],
                        colour,
                        3,
                    )

                    midpoint = (
                        (gate["p1"][0] + gate["p2"][0]) // 2,
                        (gate["p1"][1] + gate["p2"][1]) // 2,
                    )

                    cv2.putText(
                        frame,
                        gate["name"],
                        midpoint,
                        cv2.FONT_HERSHEY_SIMPLEX,
                        0.8,
                        colour,
                        2,
                    )

                # Draw vehicle boxes and IDs.
                for detection in detections:
                    x1, y1, x2, y2 = detection["bbox"]
                    vehicle_id = detection["vehicle_id"]

                    cv2.rectangle(
                        frame,
                        (x1, y1),
                        (x2, y2),
                        (0, 255, 0),
                        2,
                    )

                    cv2.putText(
                        frame,
                        f"{vehicle_id} | {detection['vehicle_type']}",
                        (x1, max(20, y1 - 8)),
                        cv2.FONT_HERSHEY_SIMPLEX,
                        0.55,
                        (0, 255, 0),
                        2,
                    )

                progress = f"Frame {frame_number}/{total_frames}"
                cv2.putText(
                    frame,
                    progress,
                    (12, height - 18),
                    cv2.FONT_HERSHEY_SIMPLEX,
                    0.6,
                    (255, 255, 255),
                    2,
                )

                cv2.imshow("Vehicle Forensic Analysis System", frame)

                if frame_number % 100 == 0:
                    print(
                        f"Processed {frame_number}/{total_frames} frames"
                    )

                key = cv2.waitKey(1) & 0xFF
                if key in (ord("q"), 27):
                    print("\nStopped by user.")
                    break

        except KeyboardInterrupt:
            print("\nProcessing interrupted.")

        finally:
            capture.release()
            cv2.destroyAllWindows()

            # Save Excel BEFORE the temporary snapshot files are deleted.
            try:
                workbook.save(output_path)
                print(f"\nExcel file saved successfully:\n{output_path}")
            except PermissionError:
                print(
                    "\nERROR: Close vehicle_results.xlsx in Excel and run "
                    "the program again."
                )
            except Exception as error:
                print(f"\nERROR saving Excel: {error}")

    print(f"\nUnique tracking IDs created: {next_id - 1}")
    print(f"Excel records: {worksheet.max_row - 1}")
    print("Snapshots are embedded inside the Excel file.")
    print("No separate snapshot folder is needed.")
    print("\nAnalysis complete.")


if __name__ == "__main__":
    main()
